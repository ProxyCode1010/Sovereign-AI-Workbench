from pptx import Presentation
from pptx.util import Inches
from config import OUTPUT_DIR
import os

def create_slide_deck(filename: str, slides: list[dict]) -> str:
    """slides: [{"title": ..., "bullets": [...]}]"""
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for s in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s["title"]
        body = slide.placeholders[1].text_frame
        for i, bullet in enumerate(s["bullets"]):
            if i == 0:
                body.text = bullet
            else:
                body.add_paragraph().text = bullet
    path = os.path.join(OUTPUT_DIR, filename)
    prs.save(path)
    return path